#!/usr/bin/env python3
"""Read-only source extraction (requires openpyxl). Never writes the workbook."""
import json,re,unicodedata,hashlib,sys
from pathlib import Path
from openpyxl import load_workbook
SOURCE=Path(sys.argv[1]);ROOT=Path(__file__).resolve().parents[1]
clean=lambda s:re.sub(r'\s+',' ',s).strip()
fold=lambda s:''.join(c for c in unicodedata.normalize('NFD',s) if unicodedata.category(c)!='Mn').casefold()
entries={};excluded=[];rows=0
for sheet in load_workbook(SOURCE,data_only=True):
 for row in range(5,sheet.max_row+1):
  if not re.search(r'\b(?:PETIT|DAHMANI|DAMANI)\b',str(sheet[f'R{row}'].value or ''),re.I):continue
  rows+=1
  for section,col in [('itc','V'),('catenaire','W')]:
   cell=sheet[f'{col}{row}'];value=cell.value
   if value is None:continue
   if not isinstance(value,str):
    excluded.append({'source':f'{sheet.title}!{cell.coordinate}','reason':'Valeur non textuelle ; référence technique non interprétable.'});continue
   for raw in value.splitlines():
    label=clean(raw)
    if not label:continue
    if fold(label).startswith('coactivite'):
     excluded.append({'source':f'{sheet.title}!{cell.coordinate}','text':label,'reason':'Note de coactivité, pas une référence technique.'});continue
    kind='ZEP' if section=='itc' else 'SEL' if re.match(r'^(sel\b|\d+bis\b)',fold(label)) else 'Secteur'
    reference=re.sub(r'^(?:s[eé]l|secteur|zep)\s+','',label,flags=re.I);m=re.search(r'\b(?:Sr|Secteur)\s+.+',label,re.I) if kind=='SEL' else None
    key=(section,fold(label));src={'sheet':sheet.title.strip(),'row':row,'cell':cell.coordinate}
    e=entries.setdefault(key,{'id':hashlib.sha256((section+'|'+fold(label)).encode()).hexdigest()[:24],'section_key':section,'type':kind,'label':label,'reference':reference,'sector':m.group(0) if m else '', 'sites':[],'sources':[]})
    site=clean(str(sheet[f'S{row}'].value or ''))
    if site and site not in e['sites']:e['sites'].append(site)
    if src not in e['sources']:e['sources'].append(src)
data={'source_file':SOURCE.name,'scope':'Toutes les feuilles du classeur, y compris les semaines antérieures/ultérieures et les week-ends ; demandeurs Petit / Dahmani uniquement. Ni horaires ni numéros de téléphone importés.','source_sha256':hashlib.sha256(SOURCE.read_bytes()).hexdigest(),'matched_rows':rows,'entries':sorted(entries.values(),key=lambda e:(e['section_key'],e['type'],fold(e['label']))),'excluded_notes':excluded}
(ROOT/'data/cr-references-v156.json').write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n');print('Source rows:',rows,'References:',len(entries))
